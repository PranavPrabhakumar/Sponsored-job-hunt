"""Make the layout-cv.* fixtures: CVs that put the job title and the organisation on one line, as many templates do.

layout-cv.docx         title<tab>organisation<tab>dates, bullets from Word's "List Bullet" style
layout-cv.pdf          the same CV exported to PDF with a bold title, the organisation after a space and right-aligned dates
layout-cv-browser.pdf  a CV printed from a web page: no bullet signs in the text, dates pushed right with flexbox

All names and details are made up. Needs python-docx, LibreOffice (soffice) and Playwright with Chromium.
Run: python3 tests/fixtures/make_layout_cvs.py
"""
import shutil
import subprocess
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt
from playwright.sync_api import sync_playwright

HERE = Path(__file__).parent
PROFILE = ("Analyst with an MSc in Business Analytics and experience using Python and SQL to apply statistical models to real-",
           "world problems, from churn prediction to demand forecasting.")
JOBS = [
    ("Data Analytics Intern", "Brightwave (Remote)", "03/2025 – 05/2025",
     ["Interpreted sales datasets, delivering insights that informed pricing decisions",
      "Designed and deployed 8 interactive dashboards in Power BI"]),
    ("Business Analyst Intern", "Kestrel Data Solutions Pvt Ltd (Chennai, India)", "01/2023 – 06/2023",
     ["Built weekly SQL reports for 4 teams", "Automated data cleaning in Python, saving 5 hours a week"]),
]
EDUCATION = [("MSc Business Analytics (Merit)", "University of Leicester", "09/2023 – 09/2024"),
             ("BE Electronics and Communication Engineering", "PSG College of Technology", "06/2018 – 05/2022")]


def heading(doc, text):
    doc.add_paragraph().add_run(text).bold = True


def word_cv(layout):
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)
    r = doc.add_paragraph().add_run("JORDAN LEE"); r.bold = True; r.font.size = Pt(18)
    doc.add_paragraph("Leicester, UK | jordan.lee@example.com | 07700 900456 | linkedin.com/in/jordan-lee-demo")
    heading(doc, "PROFILE")
    p = doc.add_paragraph(); p.add_run(PROFILE[0]); p.add_run().add_break(WD_BREAK.LINE); p.add_run(PROFILE[1])
    heading(doc, "SKILLS")
    doc.add_paragraph("Python, SQL, Power BI, Tableau, Excel, scikit-learn")
    heading(doc, "EXPERIENCE")
    for title, org, dates, points in JOBS:
        p = doc.add_paragraph()
        if layout == "tabs":   # Title<tab>Organisation<tab>Dates
            p.paragraph_format.tab_stops.add_tab_stop(Inches(2.6), WD_TAB_ALIGNMENT.LEFT)
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.0), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(title).bold = True; p.add_run("\t" + org); p.add_run("\t" + dates)
        else:                  # **Title** Organisation<tab>Dates
            p.paragraph_format.tab_stops.add_tab_stop(Inches(6.0), WD_TAB_ALIGNMENT.RIGHT)
            p.add_run(title).bold = True; p.add_run(" " + org).italic = True; p.add_run("\t" + dates).font.size = Pt(9)
        for point in points:
            doc.add_paragraph(point, style="List Bullet")
    heading(doc, "EDUCATION")
    for degree, school, dates in EDUCATION:
        p = doc.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Inches(6.0), WD_TAB_ALIGNMENT.RIGHT)
        p.add_run(degree).bold = True; p.add_run(" " + school); p.add_run("\t" + dates)
    heading(doc, "LANGUAGES")
    doc.add_paragraph("English (Fluent), Tamil (Native)")
    return doc


BROWSER_CV = """<!doctype html><html><head><meta charset="utf-8"><style>
body{font-family:Arial,sans-serif;font-size:11px;width:640px;margin:30px}.row{display:flex;justify-content:space-between;align-items:center}
.d{font-size:9px;color:#555}h2{font-size:12px;margin:12px 0 4px;letter-spacing:.05em}ul{margin:2px 0 6px 18px;padding:0}p{width:430px}
</style></head><body>
<div style="font-size:20px;font-weight:bold">JORDAN LEE</div><div>Leicester, UK | jordan.lee@example.com | 07700 900456</div>
<h2>PROFILE</h2><p>{profile}</p><h2>EXPERIENCE</h2>{jobs}<h2>EDUCATION</h2>{education}</body></html>"""


def browser_cv():
    jobs = "".join(f'<div class="row"><span><b>{t}</b> {o}</span><span class="d">{d}</span></div><ul>{"".join(f"<li>{x}</li>" for x in pts)}</ul>'
                   for t, o, d, pts in JOBS)
    edu = "".join(f'<div class="row"><span><b>{g}</b> {s}</span><span class="d">{d}</span></div>' for g, s, d in EDUCATION)
    return BROWSER_CV.replace("{profile}", "".join(PROFILE)).replace("{jobs}", jobs).replace("{education}", edu)


def main():
    tmp = Path(tempfile.mkdtemp())
    word_cv("tabs").save(HERE / "layout-cv.docx")
    word_cv("space").save(tmp / "layout-cv.docx")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(tmp), str(tmp / "layout-cv.docx")],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    shutil.copy(tmp / "layout-cv.pdf", HERE / "layout-cv.pdf")
    (tmp / "browser.html").write_text(browser_cv())
    with sync_playwright() as p:
        b = p.chromium.launch(); page = b.new_page()
        page.goto((tmp / "browser.html").as_uri()); page.pdf(path=str(HERE / "layout-cv-browser.pdf"), format="A4")
        b.close()
    shutil.rmtree(tmp, ignore_errors=True)
    print("wrote layout-cv.docx, layout-cv.pdf, layout-cv-browser.pdf")


if __name__ == "__main__":
    main()
